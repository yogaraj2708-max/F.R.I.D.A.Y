"""
F.R.I.D.A.Y. 3.0 — Structured Document Editor
Performs targeted OpenXML mutations on specific paragraphs, table cells, or sections
while strictly preserving all unrelated content, runs, styles, and tables.
Enforces the anti-hallucination and clarification rules.
"""

import os
import re
import shutil
import logging
from typing import Dict, Any, Optional, List, Tuple

import docx
from docx.document import Document as DocxDocumentType

from friday_core.document.models import (
    DocEditOperation,
    DocEditRequest,
    DocEditResult,
    RelevantRegion
)
from friday_core.document.parser import DocxStructuralParser

logger = logging.getLogger("FRIDAY.DocumentEditor")


class DocxStructuredEditor:
    """
    Executes surgical OpenXML modifications on targeted DOCX structures.
    """

    @classmethod
    def execute_edit(
        cls,
        file_path: str,
        relevant_region: RelevantRegion,
        operation: DocEditOperation,
        content: Optional[str] = None,
        output_path: Optional[str] = None
    ) -> Tuple[bool, str, Optional[str], Dict[str, Any]]:
        """
        Executes targeted editing on the document based on relevant_region.
        Returns: (success: bool, message: str, final_output_path: Optional[str], details: Dict[str, Any])
        """
        if not os.path.exists(file_path):
            return False, f"Source DOCX file not found: {file_path}", None, {}

        # 1. Ambiguity Guard: Never guess or invent content when ambiguous
        if relevant_region.is_ambiguous:
            return False, f"Clarification Required: {relevant_region.ambiguity_reason}", None, {"ambiguous": True}

        # 2. Content Requirement Check for FILL / REPLACE / INSERT / APPEND
        needs_content = operation in [
            DocEditOperation.FILL,
            DocEditOperation.REPLACE,
            DocEditOperation.INSERT,
            DocEditOperation.APPEND,
            DocEditOperation.EDIT
        ]

        if needs_content and (not content or not content.strip()):
            targets_desc = ", ".join(relevant_region.matched_labels) if relevant_region.matched_labels else "the selected region"
            return (
                False,
                f"Clarification Required: I identified {targets_desc} in '{os.path.basename(file_path)}', "
                f"but what specific content or values should I {operation.value} into this section, Boss?",
                None,
                {"ambiguous": True, "needs_content": True}
            )

        # 3. Open document using python-docx
        target_out = output_path or file_path
        # Create temporary working copy to protect original against mid-write corruption
        backup_path = f"{file_path}.friday_bak"
        try:
            shutil.copy2(file_path, backup_path)
            doc = docx.Document(file_path)

            modified_targets = []
            original_texts = []

            # 4. Execute targeted modification
            if relevant_region.region_type in ("paragraph", "section") and relevant_region.paragraph_indices:
                for p_idx in relevant_region.paragraph_indices:
                    if 0 <= p_idx < len(doc.paragraphs):
                        p = doc.paragraphs[p_idx]
                        orig_text = p.text
                        original_texts.append(orig_text)

                        if operation == DocEditOperation.FILL:
                            # If paragraph has a label (e.g. "12. " or "Item 12: "), preserve the label and fill the rest
                            label_match = re.match(r"^(\s*\d+[\.\)\:\-]\s*|\s*(?:item|section|clause|question|point|field|step)\s*\d+[\.\:\-]?\s*)", orig_text, re.IGNORECASE)
                            prefix = label_match.group(1) if label_match else ""
                            # Update text cleanly while preserving existing paragraph format
                            p.text = f"{prefix}{content.strip()}"
                            modified_targets.append(f"Paragraph {p_idx} ('{orig_text[:20]}...' -> '{p.text[:20]}...')")

                        elif operation in (DocEditOperation.REPLACE, DocEditOperation.EDIT):
                            p.text = content.strip()
                            modified_targets.append(f"Paragraph {p_idx} replaced")

                        elif operation == DocEditOperation.APPEND:
                            p.text = f"{p.text} {content.strip()}"
                            modified_targets.append(f"Paragraph {p_idx} appended")

                        elif operation == DocEditOperation.DELETE:
                            p.text = ""
                            modified_targets.append(f"Paragraph {p_idx} cleared")

            elif relevant_region.region_type == "table" and relevant_region.table_cell_targets:
                for target in relevant_region.table_cell_targets:
                    t_idx = target["table"]
                    r_idx = target["row"]
                    c_idx = target["col"]
                    if 0 <= t_idx < len(doc.tables):
                        table = doc.tables[t_idx]
                        if 0 <= r_idx < len(table.rows):
                            row = table.rows[r_idx]
                            if 0 <= c_idx < len(row.cells):
                                cell = row.cells[c_idx]
                                orig_text = cell.text
                                original_texts.append(orig_text)
                                cell.text = content.strip()
                                modified_targets.append(f"Table {t_idx+1} [R{r_idx}, C{c_idx}] ('{orig_text}' -> '{cell.text}')")

            else:
                if os.path.exists(backup_path):
                    os.remove(backup_path)
                return False, f"Unsupported or empty target region: {relevant_region.region_type}", None, {}

            if not modified_targets:
                if os.path.exists(backup_path):
                    os.remove(backup_path)
                return False, "Target region matched no physical elements in document.", None, {}

            # 5. Write-back to target path
            os.makedirs(os.path.dirname(os.path.abspath(target_out)), exist_ok=True)
            doc.save(target_out)

            # Cleanup backup on successful save
            if os.path.exists(backup_path):
                os.remove(backup_path)

            details = {
                "modified_count": len(modified_targets),
                "modified_targets": modified_targets,
                "original_texts": original_texts,
                "new_content": content
            }
            msg = f"Successfully updated {len(modified_targets)} target element(s) in '{os.path.basename(target_out)}'."
            return True, msg, target_out, details

        except Exception as e:
            # Restore backup if error occurred
            if os.path.exists(backup_path):
                shutil.copy2(backup_path, file_path)
                os.remove(backup_path)
            logger.error(f"Error during structured DOCX edit: {e}", exc_info=True)
            return False, f"DOCX edit failed with error: {e}", None, {}
