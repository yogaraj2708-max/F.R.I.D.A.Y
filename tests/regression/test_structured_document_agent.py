"""
Structured Document Agent & Zero-Trust DOCX Modification Test Suite
Validates:
1. Structural parsing of real OpenXML DOCX files (DocxStructuralParser).
2. Relevant-section detection and surgical boundary identification (12 to 13).
3. Token budget enforcement (< 1,000 tokens context, preventing context overflow).
4. Ambiguity detection and non-guessing clarification behavior.
5. Missing-content clarification (no hallucinations).
6. Surgical OpenXML mutation and paragraph preservation.
7. Independent post-write verification (DocxIndependentVerifier).
8. End-to-end integration via FridayBrain.execute_smart_skill.
9. Failure-injection resilience.
"""

import os
import unittest
import asyncio
from unittest.mock import MagicMock
from docx import Document

from friday_core.document.models import DocEditOperation, DocEditResult
from friday_core.document.parser import DocxStructuralParser
from friday_core.document.editor import DocxStructuredEditor
from friday_core.document.verifier import DocxIndependentVerifier
from friday_core.document.agent import StructuredDocumentAgent
from friday_ui.core.engine import FridayBrain, FridaySignals


class TestStructuredDocumentAgent(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.test_dir = os.path.join(os.getcwd(), "scratch", "test_docx_workspace")
        os.makedirs(cls.test_dir, exist_ok=True)

        # 1. Create a structured test DOCX file with 25 numbered items and a table
        cls.sample_docx = os.path.join(cls.test_dir, "sample_contract.docx")
        doc = Document()
        doc.add_heading("Tactical Operations Framework", level=1)
        doc.add_paragraph("This framework governs system execution and deployment protocols.")
        doc.add_heading("Section 1: Operating Directives", level=2)

        for i in range(1, 26):
            doc.add_paragraph(f"{i}. Execution protocol step {i}: Standard standby operational state.")

        doc.add_heading("Section 2: Sign-Off Matrix", level=2)
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "Role"
        table.cell(0, 1).text = "Status"
        table.cell(1, 0).text = "Commander"
        table.cell(1, 1).text = "Pending"
        table.cell(2, 0).text = "Auditor"
        table.cell(2, 1).text = "Pending"

        doc.save(cls.sample_docx)

        # 2. Create an ambiguous DOCX file where Item 12 appears in two separate sections
        cls.ambiguous_docx = os.path.join(cls.test_dir, "ambiguous_contract.docx")
        adoc = Document()
        adoc.add_heading("Framework Alpha", level=1)
        adoc.add_paragraph("12. Alpha item twelve directive.")
        adoc.add_heading("Framework Beta", level=1)
        adoc.add_paragraph("12. Beta item twelve directive.")
        adoc.save(cls.ambiguous_docx)

        # 3. Create a large DOCX file to test strict token bounding
        cls.large_docx = os.path.join(cls.test_dir, "large_corporate_policy.docx")
        ldoc = Document()
        ldoc.add_heading("Massive Enterprise System Policy", level=1)
        for i in range(1, 150):
            ldoc.add_paragraph(
                f"{i}. Clause {i}: The quick brown fox jumps over the lazy dog. "
                f"Continuous streaming telemetry logs and persistent zero-trust architecture parameters "
                f"shall be adhered to by all tactical agents operating under Protocol {i}."
            )
        ldoc.save(cls.large_docx)

        # Initialize FridayBrain
        cls.signals = FridaySignals()
        cls.mock_tts = MagicMock()
        cls.mock_tts.is_available.return_value = False
        cls.mock_tts.cancel_event.is_set.return_value = False
        cls.mock_tts.stop_speaking = MagicMock()
        cls.brain = FridayBrain(cls.signals, cls.mock_tts)

    def test_01_structural_parser(self):
        """Verifies structural indexing, paragraph maps, and headings extraction."""
        doc_map = DocxStructuralParser.parse(self.sample_docx)
        self.assertEqual(doc_map.file_name, "sample_contract.docx")
        self.assertGreater(doc_map.total_paragraphs, 25)
        self.assertEqual(len(doc_map.tables), 1)
        self.assertGreaterEqual(len(doc_map.headings), 3)

        # Verify items 12 and 13 exist in item index map
        self.assertIn("12", doc_map.item_index_map)
        self.assertIn("13", doc_map.item_index_map)
        self.assertIn("Execution protocol step 12", doc_map.item_index_map["12"][0].text)
        self.assertIn("Execution protocol step 13", doc_map.item_index_map["13"][0].text)

    def test_02_relevant_section_detection(self):
        """Verifies targeting of specific numbered ranges without extracting entire document."""
        doc_map = DocxStructuralParser.parse(self.sample_docx)
        region = DocxStructuralParser.find_relevant_region(doc_map, "fill this from 12 to 13 with Approved")

        self.assertFalse(region.is_ambiguous)
        self.assertIn("Item 12", region.matched_labels)
        self.assertIn("Item 13", region.matched_labels)
        self.assertEqual(len(region.matched_paragraphs), 2)
        self.assertIn("Execution protocol step 12", region.context_text)
        self.assertIn("Execution protocol step 13", region.context_text)
        # Ensure earlier and later items outside context window are NOT in the matched region
        self.assertNotIn("Execution protocol step 10", region.context_text)
        self.assertNotIn("Execution protocol step 15", region.context_text)
        # Ensure only 12 and 13 are marked TARGET
        self.assertIn("[P14] (TARGET)", region.context_text)
        self.assertIn("[P15] (TARGET)", region.context_text)
        self.assertNotIn("[P13] (TARGET)", region.context_text)
        self.assertNotIn("[P16] (TARGET)", region.context_text)

    def test_03_token_budget_enforcement(self):
        """Verifies hard context token bounding (< 1,000 tokens) on large documents."""
        doc_map = DocxStructuralParser.parse(self.large_docx)
        region = DocxStructuralParser.find_relevant_region(doc_map, "fill this from 12 to 13")

        self.assertFalse(region.is_ambiguous)
        # 150 paragraphs would take ~6,000 tokens, but targeted region must be strictly under 1,000
        self.assertLess(region.context_token_estimate, 500)
        self.assertLess(len(region.context_text), 1500)

    def test_04_ambiguity_detection(self):
        """Verifies that duplicate item references trigger clarification rather than random guessing."""
        doc_map = DocxStructuralParser.parse(self.ambiguous_docx)
        region = DocxStructuralParser.find_relevant_region(doc_map, "fill item 12 with Done")

        self.assertTrue(region.is_ambiguous)
        self.assertIn("multiple sections", region.ambiguity_reason.lower())

    def test_05_clarification_on_missing_content(self):
        """Verifies that 'fill this from 12 to 13' without values asks for clarification."""
        agent = StructuredDocumentAgent()
        result = asyncio.run(agent.execute_task(
            file_path=self.sample_docx,
            user_directive="fill this from 12 to 13"
        ))

        self.assertFalse(result.success)
        self.assertIn("Clarification Required", result.message)
        self.assertIn("Item 12", result.message)
        self.assertIn("Item 13", result.message)

    def test_06_surgical_edit_and_independent_verification(self):
        """Verifies surgical modification, write-back, and independent verification."""
        out_docx = os.path.join(self.test_dir, "modified_contract.docx")
        agent = StructuredDocumentAgent()

        new_val = "Verified and Approved by Mission Control"
        result = asyncio.run(agent.execute_task(
            file_path=self.sample_docx,
            user_directive=f'fill this from 12 to 13 with "{new_val}"',
            output_path=out_docx
        ))

        self.assertTrue(result.success, f"Document edit failed: {result.message}")
        self.assertIn("Successfully updated", result.message)
        self.assertIn("Independent Verification: Passed", result.message)

        # Independent physical check: Read the output file directly
        self.assertTrue(os.path.exists(out_docx))
        self.assertGreater(os.path.getsize(out_docx), 0)

        verified_doc = Document(out_docx)
        all_text = [p.text for p in verified_doc.paragraphs]

        # Assert items 12 and 13 were updated
        item_12_text = [t for t in all_text if t.startswith("12.")][0]
        item_13_text = [t for t in all_text if t.startswith("13.")][0]
        self.assertIn(new_val, item_12_text)
        self.assertIn(new_val, item_13_text)

        # Assert surrounding items (11 and 14) are completely untouched
        item_11_text = [t for t in all_text if t.startswith("11.")][0]
        item_14_text = [t for t in all_text if t.startswith("14.")][0]
        self.assertIn("Execution protocol step 11: Standard standby operational state.", item_11_text)
        self.assertIn("Execution protocol step 14: Standard standby operational state.", item_14_text)

    def test_07_engine_integration_fast_path(self):
        """Verifies that FridayBrain.execute_smart_skill executes the document agent end-to-end."""
        cmd = f'in "{self.sample_docx}" fill this from 12 to 13 with "Audited by F.R.I.D.A.Y."'
        res = asyncio.run(self.brain.execute_smart_skill(cmd))

        self.assertIsNotNone(res)
        self.assertIn("Successfully updated", res)
        self.assertIn("Independent Verification: Passed", res)

    def test_08_failure_injection_nonexistent_file(self):
        """Verifies truthful failure when document does not exist (never fake success)."""
        dead_file = os.path.join(self.test_dir, "nonexistent_contract_ghost.docx")
        agent = StructuredDocumentAgent()
        result = asyncio.run(agent.execute_task(
            file_path=dead_file,
            user_directive="fill this from 12 to 13 with Done"
        ))

        self.assertFalse(result.success)
        self.assertIn("not found", result.message.lower())


if __name__ == "__main__":
    unittest.main()
